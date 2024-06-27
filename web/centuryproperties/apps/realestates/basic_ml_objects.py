import warnings
import os
from django.conf import settings
import matplotlib as mpl
import numpy as np
import pandas as pd
from django.apps import apps
from openpyxl import Workbook, load_workbook
import math
import time
import shutil
from statistics import mean
import pickle
from django.db.models import Q

import os
from django.http import HttpResponse
from django.conf import settings
import pandas as pd
import json
import time


class BaseDataProcessing(object):
    def __init__(self, dic):  # to_data_path, target_field
        # print("90002-000 BaseDataProcessing\n", dic, '\n', '-'*50)
        try:
            super(BaseDataProcessing, self).__init__(dic)  # (dic)
        except Exception as ex:
            print("Error 90002-010 \n" + str(ex),"\n", "-"*50)

        # print("90002-010 BaseDataProcessing\n", dic, '\n', '-'*50)

        self.name = 'DataProcessing'
        self.uploaded_filename = None
        # print("90003-0 PBaseDataProcessing", dic, '-'*50)

        warnings.filterwarnings(action="ignore", message="^internal gelsd")
        # to make this notebook's output stable across runs

        self.RANDOM_STATE = 42
        np.random.seed(self.RANDOM_STATE)

        # To plot pretty figures
        mpl.rc('axes', labelsize=14)
        mpl.rc('xtick', labelsize=12)
        mpl.rc('ytick', labelsize=12)
        # Where to save the figures
        # --- Change to this when you start to use Django ---
        self.PROJECT_ROOT_DIR = os.path.join(settings.WEB_DIR, "data", dic["app"])
        # print(self.PROJECT_ROOT_DIR)
        # print('-'*50)
        os.makedirs(self.PROJECT_ROOT_DIR, exist_ok=True)
        self.TOPIC_ID = dic["topic_id"]  # "fundamentals"
        self.TO_DATA_PATH = os.path.join(self.PROJECT_ROOT_DIR, "datasets")
        os.makedirs(self.TO_DATA_PATH, exist_ok=True)
        self.TO_EXCEL = os.path.join(self.TO_DATA_PATH, "excel", self.TOPIC_ID)
        os.makedirs(self.TO_EXCEL, exist_ok=True)
        self.TO_EXCEL_OUTPUT = os.path.join(self.TO_EXCEL, "output")
        os.makedirs(self.TO_EXCEL_OUTPUT, exist_ok=True)
        self.IMAGES_PATH = os.path.join(self.PROJECT_ROOT_DIR, "images")
        os.makedirs(self.IMAGES_PATH, exist_ok=True)
        self.MODELS_PATH = os.path.join(self.PROJECT_ROOT_DIR, "models")
        os.makedirs(self.MODELS_PATH, exist_ok=True)
        self.PICKLE_PATH = os.path.join(self.PROJECT_ROOT_DIR, "pickle")
        os.makedirs(self.PICKLE_PATH, exist_ok=True)
        self.target_folder = None

        # self.TARGET_FIELD = target_field
        # self.DATA = None
        # self.TRAIN = None
        # self.TEST = None
        # self.TRAIN_TARGET = None
        # self.TRAIN_DATA = None
        # self.TEST_TARGET = None
        # self.TEST_DATA = None
        # self.train_data = None
        # self.test_data = None
        # self.num_attribs = None
        # self.extra_attribs = None
        # self.model = None
        # self.HASH = hashlib.md5
        # self.PIPELINE = None

        # print('-'*50)
        # print('9010 - End constructor parent')
        # print('-'*50)

    def upload_file(self, dic):
        # print("9002 BaseDataProcessing upload_file:")
        # print(dic)
        # print("upload_file:")

        upload_file_ = dic["request"].FILES['drive_file']
        result = {}
        # We can extend and add another property: data_folder
        # like topic_id. But, we need to add this property to: params in the core view
        # and use it here.
        # for example: if data_folder=excel we choose self.TO_EXCEL

        # print("target_folder = self.TO_"+dic["folder_type"].upper())
        self.target_folder = eval("self.TO_" + dic["folder_type"].upper())
        print('dic["folder_type"].upper()-== ', dic["folder_type"].upper())
        print('====self.target_folder======', self.target_folder)
        filename = dic["request"].POST['filename']
        print('101', filename)
        self.uploaded_filename = filename
        file_path = os.path.join(self.target_folder, filename)
        with open(file_path, 'wb+') as destination:
            for c in upload_file_.chunks():
                destination.write(c)

        # print("9888-8 Uploaded\n", "-" * 30)
        result['file_path'] = file_path
        return result

    def get_general_data(self, dic):
        print("9012-9012 BaseDataProcessing get_general_data:\n", dic, "\n", "="*50, "\n")
        app_ = dic["app"]
        result = {}
        n__=0
        for k in dic["dimensions"]:
            # print("\n", k, "\n", "="*50)
            dic_ = dic["dimensions"][k]
            s = k + ' = {}'
            try:
                exec(s)
                model_name_ = dic_["model"]
                model_ = apps.get_model(app_label=app_, model_name=model_name_)
                p_key_field_name = model_._meta.pk.name
                # print("p_key_field_name", p_key_field_name)
                if p_key_field_name == "user":
                    p_key_field_name +="_id"
                # if p_key_field_name != "id":
                #     p_key_field_name +="_id"
                try:
                    s_ = ""
                    filters = dic_["filters"]
                    if filters:
                        for j in filters:
                            s_ += ".filter(" + j + "=" + str(filters[j]) + ")"
                            # print(s_)
                except Exception as ex:
                    pass
                    # print("Error 9012-9012-1: "+str(ex))

                sq = 'model_.objects'+ s_ +'.all()'
                # print("sq1:  ", sq)
                try:
                    qs = eval(sq)
                    # df_us = pd.DataFrame(list(qs.values()))
                    # print(df_us)
                    # print(qs)
                except Exception as ex:
                    print("er555-5 qs\n", str(ex))
                for r in qs:
                    try:
                        k_ = eval("r."+p_key_field_name)
                    except Exception as ex:
                        print(ex)

                    s = k + '["'+str(k_)+'"] = r.' + dic["dimensions"][k]["field_name"]
                    # print(s)
                    exec(s)
                # print(eval(k))
            except Exception as ex:
                print("err 1000: " + str(ex))

            # print('result[k] = ' + k)
            exec('result[k] = ' + k)
        # print(result)
        return result

    def clean_name(self, name):
        name_ = name.replace("-", "").replace(" ", "").replace("_", "")
        return name_

    def general_data(self, dic):
        action_ = dic['action']
        app_ = dic['app']
        group_ = dic['group']
        data_name_ = dic['data_name']
        general_data_model = apps.get_model(app_label="core", model_name="generaldata")
        if action_ == "set":
            data_json_ = dic['data_json']
            obj, is_created = general_data_model.objects.get_or_create(app=app_, group=group_, data_name=data_name_)
            obj.data_json = data_json_
            obj.save()
        else:
            obj = general_data_model.objects.get(app=app_, group=group_, data_name=data_name_)
            return obj.data_json

    def get_next_number(self, dic):
        app_ = dic['app']
        group_ = "general"
        data_name_ = "number"
        general_data_model = apps.get_model(app_label="core", model_name="generaldata")
        data_json_ = {"number": 1}
        return_number = 1
        obj, is_created = general_data_model.objects.get_or_create(app=app_, group=group_, data_name=data_name_)
        if is_created:
            obj.data_json = data_json_
            obj.save()
        else:
            # obj = general_data_model.objects.get(app=app_, group=group_, data_name=data_name_)
            data_json_ = obj.data_json
            return_number = data_json_["number"] + 1
            data_json_["number"] = return_number
            obj.data_json = data_json_
            obj.save()
        return return_number


class BasePotentialAlgo(object):
    def __init__(self, dic):  # to_data_path, target_field
        # print('-'*50, '\n', "90003-000 BasePotentialAlgo", dic, '\n', '-'*50)
        try:
            super(BasePotentialAlgo, self).__init__(dic)
        except Exception as ex:
            print("Error 90003-010 " + str(ex))

        # print("90003-020 PotentialAlgo", dic, '\n', '-'*50)

        self.entity_name = dic["entity_name"]
        self.second_time_save = ''
        self.to_save = []
        self.to_save_all = []
        self.save_to_file = None
        self.df_index = None
        app_ = dic["app"]
        measure_model_name_ = dic["measure_model"]
        model_measure_dim = apps.get_model(app_label=app_, model_name=measure_model_name_)
        self.measures_name = pd.DataFrame(model_measure_dim.objects.all().values('id', 'measure_name'))
        # print('self.measures_name\n', self.measures_name)
        try:
            self.entity_name = dic["entity_name"]
            # print(self.entity_name)
            entity_model_name_ = dic["entity_model"]
            model_entity_dim = apps.get_model(app_label=app_, model_name=entity_model_name_)
            # print(model_measure_dim)
            self.entities_name = pd.DataFrame(model_entity_dim.objects.all().values('id', self.entity_name+'_name'))
            # print("self.entities_name\n",self.entities_name)

            # need to delete the next line
            # self.countries_name = pd.DataFrame(model_entity_dim.objects.all().values('id', 'country_name'))
        except Exception as ex:
            print("Error 9001-222-22: "+str(ex))
        try:
            self.total_variables = dic["total_variables"]
            self.multiply_two_variables = dic["multiply_two_variables"]
            self.do_not_include_measures_objs = dic["do_not_include_measures_objs"]
            self.do_not_include_groups = dic["do_not_include_groups"]
        except Exception as ex:
            pass
        # print(self.measures_name)
        self.options = ["mm", "mx", "xm", "xx"]
        self.is_calculate_min_max = None

    def process_algo(self, dic):
        print("90011-111 process_algo: \n", dic, "\n", "="*50)
        def get_culomns_names(self, l):
            cs = []
            for c in l:
                k_ = str(self.measures_name[self.measures_name["id"] == c]["measure_name"]).split("    ")[1].split("\n")[0]
                cs.append(k_)
            return cs
        # print("9015-1 BasePotentialAlgo process_algo\n", dic)
        app_ = dic["app"]
        fact_model_name_ = dic["fact_model"]
        dependent_group = dic["dependent_group"]
        # print('dependent_group', dependent_group)
        min_max_model_name_ = dic["min_max_model"]
        measure_group_model_name_ = dic["measure_group_model"]
        model_fact = apps.get_model(app_label=app_, model_name=fact_model_name_)
        # try:
        #     self.is_calculate_min_max = eval(dic["is_calculate_min_max"])
        #     # print(self.is_calculate_min_max)
        # except Exception as ex:
        #     print("Error 9016: \n"+str(ex))
        # if not self.is_calculate_min_max:
        model_min_max = apps.get_model(app_label=app_, model_name=min_max_model_name_)
        # else:
        #     self.calculate_min_max_cuts(dic)
        model_measure_group = apps.get_model(app_label=app_, model_name=measure_group_model_name_)
        # print(2001)
        # print("90060-10 PotentialAlgo: \n", "="*50)
        wb2 = None
        groups = model_measure_group.objects.all()
        nn__ = 0
        sign_n1 = pd.DataFrame([[0, 0, 0, 0]])
        sign_n2 = pd.DataFrame([[0, 0, 0, 0]])
        sign_n1.columns = self.options
        sign_n2.columns = self.options
        similarity_n1 = pd.DataFrame([[0, 0, 0, 0]])
        similarity_n2 = pd.DataFrame([[0, 0, 0, 0]])
        contribution_n1 = pd.DataFrame([[0, 0, 0, 0]])
        adj_contribution_n1 = pd.DataFrame([[0, 0, 0, 0]])
        contribution_n2 = pd.DataFrame([[0, 0, 0, 0]])
        adj_contribution_n2 = pd.DataFrame([[0, 0, 0, 0]])
        # relimp_n1 = pd.DataFrame([[0, 0, 0, 0]])
        # relimp_n2 = pd.DataFrame([[0, 0, 0, 0]])
        similarity_n1.columns = self.options
        similarity_n2.columns = self.options
        group_d = ""
        # print(2002)
        ll_dfs = self.pre_process_data(dic)
        # print(ll_dfs)

        lll_groups = []
        for k in ll_dfs:
            group = k #.group_name
            # print("="*50)
            # print(group)
            # print("="*50)
            try:
                self.save_to_file = os.path.join(self.TO_EXCEL_OUTPUT, str(dic["time_dim_value"]) + "_" + group + "_o.xlsx")
                self.to_save = []
                # print("file_path\n", self.save_to_file, "\n", "="*50)
                s = ""
                # for v in dic["axes"]:
                #     s += "'" + v + "',"
                # s += "'" + dic["value"] + "'"
                # # print(dic["time_dim_value"])
                # qs = model_fact.objects.filter(measure_dim__measure_group_dim__group_name=group,
                #                                time_dim_id=dic["time_dim_value"]).all()
                # s = "pd.DataFrame(list(qs.values(" + s + ")))"
                # df = eval(s)
                # if df.shape[0] == 0:
                #     continue
                # # print("50001-21")
                # df = df.pivot(index="country_dim", columns='measure_dim', values='amount')
                # # print(df)
                # # print("50001-22")
                # self.df_index = df.index
                # df_columns = df.columns
                # # print(df)
                # # print("50001-22-1")
                # df_ = self.add_entity_to_df(df)
                # self.to_save.append((df_.copy(), 'Data'))
                # #
                # if not self.is_calculate_min_max:

                df = ll_dfs[group]
                df_columns = df.columns
                self.df_index = df.index
                # print(self.df_index)
                df_ = self.add_entity_to_df(df)
                self.to_save.append((df_.copy(), 'Data'))

                qs_mm = model_min_max.objects.filter(measure_dim__measure_group_dim__group_name=group,
                                                     time_dim_id=dic["time_dim_value"]).all()
                df_mm = pd.DataFrame(list(qs_mm.values('measure_dim', 'min', 'max')))
                first_row = pd.DataFrame(df_mm.T.loc["min"]).T.reset_index().drop(['index'], axis=1)
                second_row = pd.DataFrame(df_mm.T.loc["max"]).T.reset_index().drop(['index'], axis=1)

                first_row.columns = df_columns
                second_row.columns = df_columns
                diff_row = second_row.subtract(first_row, fill_value=None)
                diff_row.columns = df_columns

                # print("first_row,second_row,diff_row\n", first_row, "\n", second_row, "\n", diff_row)

                df = ll_dfs[group]
                df_columns = df.columns
                # print("df.columns, df_columns\n", df.columns, df_columns)
                # print("df\n", df)

                df_n1 = df.copy()
                try:
                    df_n1 = df_n1.astype(float)
                except Exception as ex:
                    print("1000: " + str(ex))
                for i, r in df_n1.iterrows():
                    for j in df_columns:
                        try:
                            z = (r[j] - first_row[j].astype(float)) / diff_row[j].astype(float)
                            df_n1.loc[i][j] = z
                        except Exception as ex:
                            print("Error i " + str(i) + " " + str(ex))
                # print("1 df_n1\n", df_n1, "\n", "="*100)

                df_n1 = df_n1.apply(pd.to_numeric, errors='coerce').round(6)
                self.add_to_save(title='Normalized-1', a=df_n1, cols=None)
                #
                # print("50001-2")
                df_n2 = df_n1.copy()
                df_n2[df_n2 < 0] = 0
                df_n2[df_n2 > 1] = 1
                self.add_to_save(title='Normalized-2', a=df_n2, cols=None)
                #
                if len(df_n1.columns) < 2:
                    df_n1["max"] = df_n1[df_n1.columns[0]]  # df_n1["Birth Rate"]
                    df_n2["max"] = df_n2[df_n2.columns[0]]  # df_n2["Birth Rate"]
                    df_1_2 = pd.merge(left=df_n1, right=df_n2, left_index=True, right_index=True)
                    # df_1_2.columns = ['min-n1', 'max-n1', 'min-n2', 'max-n2']
                    # cols = df_1_2.columns
                    # cols = cols.insert(0, 'country_name')
                    # self.add_to_save(title='min-max', a=df_1_2, cols=cols)
                elif len(df_n1.columns) < 5:
                    # print("-1"*30)
                    # print(group)
                    # print("-000"*30)
                    df_n1 = df_n1.apply(lambda x: np.sort(x), axis=1, raw=True)
                    df_n2 = df_n2.apply(lambda x: np.sort(x), axis=1, raw=True)
                    df_n1["max"] = df_n1.max(axis=1)
                    df_n1["min"] = df_n1.min(axis=1)
                    df_n2["max"] = df_n2.max(axis=1)
                    df_n2["min"] = df_n2.min(axis=1)
                    # df_n1 = df_n1.drop(df_n1_columns, axis=1)
                    # df_n2 = df_n2.drop(df_n2_columns, axis=1)
                    df_n1 = df_n1[["min", "max"]]
                    df_n2 = df_n2[["min", "max"]]
                    df_1_2 = pd.merge(left=df_n1, right=df_n2, left_index=True, right_index=True)
                    # df_1_2.columns = ['min-n1', 'max-n1', 'min-n2', 'max-n2']
                    # cols = df_1_2.columns
                    # cols = cols.insert(0, 'country_name')
                    # df_ = self.add_entity_to_df(df_1_2, cols)
                    # print(df_)
                    # self.to_save.append((df_.copy(), 'min-max'))
                else:
                    zero_list = {}
                    one_list = {}
                    # print(df_n1)
                    for i in df_n1.index:
                        n0 = []
                        n1 = []
                        for num in df_n1.loc[i]:
                            if not pd.isna(num):
                                if num <= 0:
                                    n0.append(num)
                                elif num >= 1:
                                    n1.append(num)
                        if len(n0) > 0:
                            # print("n0\n", n0)
                            n0.sort(reverse=True)
                            # print("A n0\n", n0)
                            # print("="*10)
                            zero_list[i] = n0
                        elif len(n1) > 0:
                            n1.sort()
                            one_list[i] = n1
                    # # #
                    a = df_n2.values
                    a_1 = a.copy()
                    a_1.sort(axis=1)
                    self.add_to_save(title='Sort L', a=a_1, cols=-1)
                    a_1 = self.clean_rows(a=a_1, j=1, side="L")
                    #
                    a_1m = pd.DataFrame(a_1)
                    a_1m = a_1m.apply(self.move_elements_to_right, axis=1)
                    a_1m = a_1m.apply(self.revers_elements_in_row, axis=1)
                    self.add_to_save(title='R arranged', a=a_1m, cols=-1)
                    #
                    a_1 = -1 * a_1.copy()
                    a_1.sort(axis=1)
                    self.add_to_save(title='Sort R', a=a_1, cols=-1)
                    a_1 = self.clean_rows(a=a_1, j=1, side="R")
                    a_1 = -1 * a_1.copy()
                    a_1.sort(axis=1)
                    #
                    # move the numbers to the right.
                    # a_1 = pd.DataFrame(a_1)
                    # a_1 = a_1.apply(self.revers_elements_in_row, axis=1)
                    # print(a_1)
                    #
                    self.add_to_save(title='Final R', a=a_1, cols=-1)
                    # print(a_1)
                    # a_1 = -1 * a_1
                    # a_1 = a_1.apply(self.revers_elements_in_row, axis=1)
                    # a_1.sort(axis=1)  #
                    # self.add_to_save(title='Final-1', a=a_1, cols=-1)
                    # # #
                    a_1 = pd.DataFrame(a_1)
                    a_1.dropna(how='all', axis=1, inplace=True)
                    # print("600000000-100-1")
                    #
                    a_1 = a_1.apply(self.twenty_rule, axis=1)
                    # print("600000000-100-12")
                    a_1 = a_1.apply(lambda x: np.sort(x), axis=1, raw=True)
                    # print("600000000-100-13")
                    self.add_to_save(title='Final-20-rule', a=a_1, cols=-1)

                    # print("600000000-100-2")

                    a_1 = a_1.apply(self.thirty_rule, axis=1)
                    self.add_to_save(title='Final-30-rule', a=a_1, cols=-1)
                    a_2 = a_1.copy()
                    #
                    # print("zero_list\n", zero_list)
                    ff = []
                    for j in a_1.index:
                        nn = list(a_1.loc[j])
                        # print("nn == \n", nn)
                        if min(nn) == 0:
                            # print("nn zero_list[j]\n", zero_list[j])
                            nn = [zero_list[j].pop(0) if i == 0 else i for i in nn]
                            # print("nn 0\n", nn)
                        if max(nn) == 1:
                            nn = [one_list[j].pop() if i == 1 else i for i in nn]
                            # print("nn 1\n", nn)
                        # nn.insert(0, j)
                        nn.sort()
                        ff.append(nn)
                    self.add_to_save(title='Final-30-rule-n1', a=ff, cols=-1)
                    a_1 = pd.DataFrame(ff, index=list(self.df_index))
                    #
                    df_n1 = a_1.apply(self.min_max_rule, axis=1)
                    df_n1.columns = ['min-n1', 'max-n1']
                    #
                    df_n2 = a_2.apply(self.min_max_rule, axis=1)
                    df_n2.columns = ['min-n2', 'max-n2']
                    df_1_2 = pd.merge(left=df_n1, right=df_n2, left_index=True, right_index=True)
                # print("50001-3")
                df_1_2.columns = ['min-n1', 'max-n1', 'min-n2', 'max-n2']
                cols = df_1_2.columns
                cols = cols.insert(0, self.entity_name+'_name')
                # print("50001-3-6")
                self.add_to_save(title='min-max', a=df_1_2, cols=cols)
                # self.add_to_save(title='min-max', a=df_1_2, cols=-1)
                # print("50001-3-9")
                self.save_to_excel_()

                # print("50001-3-9-1")
            except Exception as ex:
                print("Error 50001-136-1: " + str(ex))
            df_n1_ = df_n1.copy()
            df_n1_.columns = ['m-' + group, 'x-' + group]
            df_n2_ = df_n2.copy()
            df_n2_.columns = ['m-' + group, 'x-' + group]

            if group == dependent_group:
                ss_n_mm = ""
                ss_n_xm = ""
                ss_n_mx = ""
                ss_n_xx = ""
                group_d = group
                df_n1_all = df_n1_
                df_n2_all = df_n2_
            else:
                # print(group_d, group)
                group_d, group, df_n1_all, df_n2_all, df_n1_, df_n2_, sign_n1, sign_n2, similarity_n1, similarity_n2 = \
                    self.create_similarity(group_d, group, df_n1_all, df_n2_all, df_n1_, df_n2_, sign_n1, sign_n2,
                                           similarity_n1, similarity_n2)
                ss_n_mm += '"' + group_d + '-' + group + '-mm",'
                ss_n_mx += '"' + group_d + '-' + group + '-mx",'
                ss_n_xm += '"' + group_d + '-' + group + '-xm",'
                ss_n_xx += '"' + group_d + '-' + group + '-xx",'
            lll_groups.append(group)
        qs_mm = model_min_max.objects.filter(measure_dim__measure_group_dim__group_name=group,
                                             time_dim_id=dic["time_dim_value"]).all()
        df_mm = pd.DataFrame(list(qs_mm.values('measure_dim', 'min', 'max')))
        first_row = pd.DataFrame(df_mm.T.loc["min"]).T.reset_index().drop(['index'], axis=1)
        # print(first_row)
        second_row = pd.DataFrame(df_mm.T.loc["max"]).T.reset_index().drop(['index'], axis=1)

        ss_n_mm = ss_n_mm[:-1]
        ss_n_mx = ss_n_mx[:-1]
        ss_n_xm = ss_n_xm[:-1]
        ss_n_xx = ss_n_xx[:-1]

        df_n2_all_corr = df_n2_all.copy()
        # print("1 - df_n2_all_corr\n", df_n2_all_corr, "\n", ll_dfs)
        # print(df_n2_all_corr.columns)
        for o in ["m", "x"]:
            s_corr = ""
            n__ = 0
            for g in ll_dfs:
                #print(g)
                if n__ >= 1:
                    s_corr += "'"+o+"-"+g+"',"
                    #print(0, n__, s_corr)
                n__ += 1
            # print(1, s_corr)
            s_corr = s_corr[:-1]
            # print(2, s_corr)

            df_corr = eval("df_n2_all_corr[["+s_corr+"]]")
            corr_ = df_corr.corr(method='pearson')
            # print("90050-30\n", corr_,"\n", type(corr_))
            self.add_to_save_all(title='corr-' + o, a=corr_, cols=-1)
        # print("2 - ")
        for n in ["1", "2"]:
            ll = []
            for k in self.options:
                s_ = "df_n" + n + "_all['d_" + k + "']=df_n" + n + "_all[[" + eval("ss_n_" + k) + "]].min(axis=1)"
                #print(s_)
                exec(s_)
                s_ = "ll.append(1-df_n" + n + "_all[[" + eval("ss_n_" + k) + "]].min(axis=1).mean())"
                #print(s_)
                exec(s_)

            s_ = "similarity_n" + n + ".loc['SComb'] = ll"
            # print(s_, "\n", ll)

            exec(s_)
            s_ = "sign_n" + n + ".drop([0], axis=0, inplace=True)"
            # print(s_)
            exec(s_)

            self.add_to_save_all(title='sign-n' + n, a=eval("sign_n" + n), cols=-1)
            exec("similarity_n" + n + ".drop([0], axis=0, inplace=True)")
            self.add_to_save_all(title='similarity-n' + n, a=eval("similarity_n" + n), cols=-1)

        # print("90050-27\n")

        for n in ["1", "2"]:
            nn__ = 0
            llg = []
            llg_adj = []
            for k in lll_groups:
                try:
                    group = k #.group_name
                    # print("-"*10)
                    # print(group)
                    if nn__ > 0:
                        ll = []
                        for z in self.options:
                            s_ = "df_n" + n + "_all['dc_" + group + "_" + z + "'] = abs("
                            s_ += "df_n" + n + "_all['d_" + z + "'] - "
                            s_ += "df_n" + n + "_all['" + group_d + "-" + group + "-" + z + "'])"
                            exec(s_)
                            # print('ll.append(1-df_n'+n+'_all["dc_'+group+'_' + z+'"].mean())')
                            exec('ll.append(1-df_n' + n + '_all["dc_' + group + '_' + z + '"].mean())')

                        ll_adj = eval("ll*similarity_n" + n + ".loc['SComb'].T").tolist()
                        print("ll\n", ll, "\n", type(ll))
                        print("ll_adj\n", ll_adj, "\n", type(ll_adj))
                        #
                        llc = [(x - 0.7) if x > 0.7 else 0 for x in ll]
                        llg.append(llc)
                        print("llg\n", llg)
                        llc_adj = [(x - 0.7) if x > 0.7 else 0 for x in ll_adj]
                        llg_adj.append(llc_adj)
                        print("llg_adj\n", llg_adj)

                        exec("contribution_n" + n + ".loc[group] = ll")
                        exec("adj_contribution_n" + n + ".loc[group] = ll_adj")
                        # print("A   contribution_n" + n, eval("contribution_n"+n), "\n", ll)
                    else:
                        nn__ += 1
                except Exception as ex:
                    print("Error 50008-8: " + str(ex))

            print(n, "\nCC\n", llg, "\n\n", llg_adj)

            self.add_to_save_all(title="all-n" + n, a=eval("df_n" + n + "_all"), cols=-1)
            exec("contribution_n" + n + ".columns = self.options")
            exec("adj_contribution_n" + n + ".columns = self.options")
            exec("contribution_n" + n + ".drop([0], axis=0, inplace=True)")
            exec("adj_contribution_n" + n + ".drop([0], axis=0, inplace=True)")

            print("=1"*50)

            npg = np.array(llg)
            npg_adj = np.array(llg_adj)
            # print("llg", "\n", llg)

            npgs = np.sum(llg, axis=0)
            npgs_adj = np.sum(llg_adj, axis=0)

            df_relimp = pd.DataFrame(npg/npgs, index=contribution_n1.index)
            df_relimp_adj = pd.DataFrame(npg_adj/npgs_adj, index=contribution_n1.index)
            df_relimp.columns = self.options
            df_relimp_adj.columns = self.options
            # print(df_relimp)
            self.add_to_save_all(title='contribution-n' + n, a=eval("contribution_n" + n), cols=-1)
            self.add_to_save_all(title='adj_contribution-n' + n, a=eval("adj_contribution_n" + n), cols=-1)
            self.add_to_save_all(title='relimp-n' + n, a=df_relimp, cols=-1)
            self.add_to_save_all(title='adj_relimp-n' + n, a=df_relimp_adj, cols=-1)
            print("=2"*50)

        # print("90050-28\n")
        self.save_to_excel_all_(dic["time_dim_value"])
        # print("90050-29\n")

        result = {"status": "ok"}
        return result

    def get_dim_data_frame(self, dic):
        # print("90066-106-11 PotentialAlgo calculate_min_max_cuts: \n", dic, "\n", "="*50)

        # dic =  {'app': 'avi', 'filter_dim': 'time_dim', 'filter_value': 2019, 'value': 'amount',
        #         'axes': ['country_dim', 'measure_dim'],
        #         'model_name': 'worldbankfact', 'filter_amount': 1000000, 'measure_id': '19',
        #         'measure_name': 'TotalPop'}
        app_ = dic["app"]
        filter_dim = dic["filter_dim"]
        filter_value = dic["filter_value"]
        value_ = dic["value"]
        axes_ = dic["axes"]
        model_name = dic["model_name"]
        filter_amount = dic["filter_amount"]  # 1000000
        measure_id = dic["measure_id"]
        #
        return_dim = ""
        for k in axes_:
            if k not in ["measure_dim", filter_dim]:
                return_dim = k
        #
        # print(return_dim)
        model = apps.get_model(app_label=app_, model_name=model_name)

        if measure_id is None or measure_id == "":
            measure_name = dic["measure_name"]
            s = 'model.objects.filter(measure_dim__measure_name="'+measure_name+'", '
        else:
            # print(measure_id)
            s = 'model.objects.filter(measure_dim__id='+measure_id+', '
        s += filter_dim+'__id='+str(filter_value)+', amount__gte='+str(filter_amount)+').all()'
        # print("1 \ns="+s)
        qs = eval(s)
        # qs = model.objects.filter(measure_dim__measure_name="TotalPop", time_dim__id=2019, amount__gte=1).all()
        # print(qs)
        df = pd.DataFrame(list(qs.values(return_dim+"_id", value_)))
        # print(df)
        # print(df[return_dim+"_id"].tolist())
        # print(df.shape)
        if df.shape[0] > 0:
            result = {"status": "ok", "result": [df, df[return_dim+"_id"].tolist()]}
        else:
            result = {"status": "ok", "result": -1}
        # print(result)
        return result

    def convert_total_to_per_capita(self, f, df_var, df_entities):
        # print("="*100)
        df_ = df_var.reset_index()
        # print(df_var, "\n" , df_)
        df_ = df_.merge(df_entities, how='inner', left_on=self.entity_name+'_dim', right_on=self.entity_name+'_dim_id')
        # .drop(['country_dim', 'id'], axis=1)
        # print(df_)
        df_["pc"] = df_[f]/df_["amount"].astype(float)
        # print(df_, "\n", "-"*10)
        df_ = df_.drop([self.entity_name+'_dim_id', 'amount', f], axis=1)
        df_ = df_.set_index(self.entity_name+'_dim')
        #
        # print(df_, "\n", "="*50)

        return df_

    def multiply_two_variables_f(self, f__, df_var1, f__2, df_var2):
        # print("="*100)
        df_var1 = df_var1.reset_index()
        df_var2 = df_var2.reset_index()
        # print("\n", df_var1, "\n", df_var2, "\n", "="*50)

        df_ = df_var1.merge(df_var2, how='inner', left_on=self.entity_name+'_dim', right_on=self.entity_name+'_dim')
        # print(df_)
        df_["m"] = df_[f__].astype(float)*df_[f__2].astype(float)
        df_ = df_.drop([f__, f__2], axis=1)
        # print(df_)
        df_ = df_.set_index(self.entity_name+'_dim')
        # print(df_)

        return df_

    def pre_process_data(self, dic):
        # print("90033-133 pre_process_data: \n", dic, "\n", "="*50)
        app_ = dic["app"]
        fact_model_name_ = dic["fact_model"]
        year_ = str(dic["time_dim_value"])

        dic_ = {'app': app_, 'filter_dim': 'time_dim', 'filter_value': year_, 'value': 'amount',
                'axes': [self.entity_name+'_dim', 'measure_dim'],
                'model_name': fact_model_name_,
                'filter_amount': dic["population_filter_amount"], 'measure_id': None,
                'measure_name': dic["measure_name"]}

        # print("90022-122-111 pre_process_data: \n", dic_, "\n", "="*50)
        returned_ = self.get_dim_data_frame(dic_)
        # print("90011-111-1 pre_process_data: \n", returned_, "\n", "="*50)
        entity_list = returned_["result"][1]
        df_entities = returned_["result"][0]
        # print("entity_list\n", entity_list, "\n df_entities\n", df_entities)

        measure_group_model_name_ = dic["measure_group_model"]
        dependent_group = dic["dependent_group"]
        model_fact = apps.get_model(app_label=app_, model_name=fact_model_name_)
        model_measure_group = apps.get_model(app_label=app_, model_name=measure_group_model_name_)

        groups = model_measure_group.objects.filter(~Q(group_name__in=self.do_not_include_groups)).all()
        ll_groups = [dependent_group]

        # print("90-111-222-0 for k in groups")
        for k in groups:
            group = k.group_name
            if group not in ll_groups and group not in self.do_not_include_groups:
                ll_groups.append(group)
        lll_groups = []  # this will have only the groups that do not have problems (have data)
        ll_dfs = {}  # includes all the df of all groups with data
        # print("90-111-222-1")
        for k in ll_groups:
            # print("="*50, "\n", k, "\n", "="*50)
            try:
                # print("\n90-111-2-"+k, "\n", "-"*30)
                s = ""
                for v in dic["axes"]:
                    s += "'" + v + "',"
                s += "'" + dic["value"] + "'"
                s_ = 'model_fact.objects.filter(measure_dim__measure_group_dim__group_name=k, time_dim_id=year_, '+self.entity_name+'_dim_id__in=entity_list).filter(~Q(measure_dim__in=self.do_not_include_measures_objs)).all()'
                # print(s_)
                qs = eval(s_)
                # print(qs)
                s = "pd.DataFrame(list(qs.values(" + s + ")))"
                df = eval(s)
                # print(df)

                if df.shape[0] == 0:
                    continue
                lll_groups.append(k)
                df = df.pivot(index=self.entity_name+"_dim", columns='measure_dim', values='amount')
                # print("90-111-2-100\n", df)
                ll_dfs[k] = df.apply(pd.to_numeric, errors='coerce').round(6)
                for f__ in ll_dfs[k]:
                    # print("90-111-3-"+k+" "+str(f__), self.total_variables, "\n", self.multiply_two_variables)
                    # print(f__)
                    if f__ in self.total_variables:
                        # print("one")
                        ll_dfs[k][f__] = self.convert_total_to_per_capita(f__, ll_dfs[k][f__], df_entities)
                        # print("one : ", f__)
                    elif len(self.multiply_two_variables) > 0:
                        if f__ in self.multiply_two_variables["var1"]:
                            # print("two")
                            # print("two1 : ", f__, self.multiply_two_variables["var1"].index(f__))
                            idx__ = self.multiply_two_variables["var1"].index(f__)
                            f__var2 = self.multiply_two_variables["var2"][idx__]
                            f__var2_group = self.multiply_two_variables["var2_group"][idx__]
                            # print("two2: ", f__, f__var2, f__var2_group)
                            # print("two3: ", k, f__, f__var2, "\n")
                            # print("three3 : ")
                            ll_dfs[k][f__] = self.multiply_two_variables_f(f__, ll_dfs[k][f__], f__var2,
                                                                           ll_dfs[f__var2_group][f__var2])
                            # print("four4: ", k, f__, f__var2, "\n", ll_dfs[k][f__])
            except Exception as ex:
                print("Error 50661-12 PotentialAlgo calculate_min_max_cuts: \n" + str(ex), "\n", "90-111-222-2-"+k+" "+f__)
        return ll_dfs

    def calculate_min_max_cuts(self, dic):
        # print("90066-100 PotentialAlgo calculate_min_max_cuts: \n", dic, "\n", "="*50)
        app_ = dic["app"]
        f = int(dic["dim_field"])
        method = dic["method"]
        u_method = "median"
        l_method = "median"
        if method == "min":
            u_method = "min"
            l_method = "max"

        # print("90066-100-1 PotentialAlgo calculate_min_max_cuts: \n", "="*50)
        dependent_group = dic["dependent_group"]
        # print(dependent_group)
        year_ = str(dic["time_dim_value"])
        # print(year_)
        min_max_model_name_ = dic["min_max_model"]
        #
        # print("90066-100-2 PotentialAlgo calculate_min_max_cuts: \n", "="*50)
        ll_dfs = self.pre_process_data(dic)
        # print(ll_dfs)

        # result = {"status": "ok"}
        # return result
        #
        # sign_ = pd.DataFrame([[0, 0, 0, 0]])
        # sign_.columns = self.options
        # similarity_ = pd.DataFrame([[0, 0, 0, 0]])
        # similarity_.columns = self.options

        high_group_cut = [0.4, 0.35, 0.30, 0.25, 0.2, 0.15, 0.1]
        low_group_cut = [0.4, 0.35, 0.30, 0.25, 0.2, 0.15, 0.1]
        step_ = 10
        df_d = ll_dfs[dependent_group]
        # print(df_d)
        results = {}
        best_cut = {"bv": -1}

        # print("90066-100-5 PotentialAlgo calculate_min_max_cuts: \n", "="*50)

        for h in high_group_cut:
            hh = 1- h
            results[hh] = {}
            #
            # z = 5
            # steps_h = list(range(z, 100, z))
            # if h == high_group_cut[0]:
            #     steps_h.append(95)
            # print(h)
            steps_h = list(range(5, round(h * 100), 5))
            steps_h = [a/100 for a in steps_h]
            print("hhhhhhh", h, steps_h)
            #
            for ll in low_group_cut:
                results[hh][ll] = {}
                # z_ = 5
                # steps_l = list(range(z_, 100, z_))
                # if ll == low_group_cut[0]:
                #     steps_l.append(95)
                steps_l = list(range(5, round(ll * 100), 5))
                steps_l = [a/100 for a in steps_l]
                print("llllll", ll, steps_l)
                #
                # print(df_d)
                df_q = df_d.quantile([hh, ll])
                # print("df_d.shape", df_d.shape)
                # print("df_q", df_q)

                results[hh][ll][f] = {}
                # results[hh][ll][f]["max_cut"] = df_q[f].iloc[0]
                # results[hh][ll][f]["min_cut"] = df_q[f].iloc[1]
                df_hi = df_d[df_d[f] >= df_q[f].iloc[0]].index
                df_li = df_d[df_d[f] <= df_q[f].iloc[1]].index
                # print("df", "\n", df_d[df_d[f] >= df_q[f].iloc[0]], "\n\n")
                # print("df", "\n", df_d[df_d[f] >= df_q[f].iloc[0]].index, "\n\n")

                # print("df", "\n", float(df_d[df_d[f] >= df_q[f].iloc[0]][f].median()), "\n")

                # results[hh][ll][f]["max_cut"] = float(df_d[df_d[f] >= df_q[f].iloc[0]][f].median())
                # results[hh][ll][f]["min_cut"] = float(df_d[df_d[f] <= df_q[f].iloc[1]][f].median())

                for steph in steps_h:
                    for stepl in steps_l:
                        for g in ll_dfs:
                            # if g != dependent_group:
                            df = ll_dfs[g]
                            # print("\n",g, "\n",df.columns)
                            iih = []
                            iil = []
                            for j in df.index:
                                if j in df_hi:
                                    iih.append(j)
                                if j in df_li:
                                    iil.append(j)
                            dfh = df.loc[iih]
                            dfl = df.loc[iil]

                            dfh_q = dfh.quantile(steph/h)
                            dfl_q = dfl.quantile((ll-stepl)/ll)
                            # print('dfh, dfh', 'steph ', steph, 'dfh_q ', dfh_q)
                            # , '\n', 'dfl', dfl, 'stepl', stepl, 'dfl_q', dfl_q)

                            if steph not in results[hh][ll][f]:
                                results[hh][ll][f][steph] = {}
                            if stepl not in results[hh][ll][f][steph]:
                                results[hh][ll][f][steph][stepl] = {}
                                results[hh][ll][f][steph][stepl]["groups"] = {}
                            if g not in results[hh][ll][f][steph][stepl]["groups"]:
                                results[hh][ll][f][steph][stepl]["groups"][g] = {}

                            for k in df.columns:
                                # print(steph, stepl, k, dfh_q[k], dfl_q[k])
                                dfh_ik = dfh[dfh[k] >= dfh_q[k]][k]
                                dfl_ik = dfl[dfl[k] <= dfl_q[k]][k]
                                # print('dfh[dfh[k] >= dfh_q[k]]', dfh[dfh[k] >= dfh_q[k]].index, 'dfh_ik.median()', dfh_ik.median(), 'dfl_ik.median()', dfl_ik.median())
                                # print('g, k, dfh_ik.shape', g, k, dfh_ik.shape)
                                if dfh_ik.shape == 0 or dfl_ik.shape == 0:
                                    continue
                                results[hh][ll][f][steph][stepl]["groups"][g][k] = eval('{"min_cut": dfl_ik.'+l_method+'(), "max_cut": dfh_ik.'+u_method+'()}')
                        # print(results)
                        dic = {"dependent_group": dependent_group, "f":f, "steph":steph, "stepl":stepl,
                               "df_d": df_d, "ll_dfs": ll_dfs, "groups": results[hh][ll][f]}
                        bv = self.get_similarity(dic)
                        if bv > best_cut["bv"]:
                            best_cut = {"bv": bv, "hh": hh, "ll": ll, "steph":steph, "stepl":stepl, "f":f,
                                        "dependent_group": dependent_group, "year": year_,
                                        "f_groups":results[hh][ll][f][steph][stepl]["groups"]}

                            # print('best_cut\n', 'hh', hh, 'll', ll, 'results 1 steph', steph, 'stepl', stepl, "dv", bv, "\n",
                            #       best_cut)
                        print('hh', hh, 'll', ll, 'steph', round(100*(steph+hh))/100, 'stepl', round(100*(ll-stepl))/100, 'dv', bv, 'best', best_cut['bv'])

        # print("="*100)
        best_cut["df_d"] = df_d
        # print("final", best_cut, "="*100)
        #
        file_path = os.path.join(self.PICKLE_PATH, "result_"+str(f)+"_"+year_+".pkl")
        print(file_path)
        with open(file_path, 'wb') as handle:
            pickle.dump(results, handle, protocol=pickle.HIGHEST_PROTOCOL)

        file_path = os.path.join(self.PICKLE_PATH, "best_cut_"+str(f)+"_"+year_+".pkl")
        print(file_path)
        with open(file_path, 'wb') as handle:
            pickle.dump(best_cut, handle, protocol=pickle.HIGHEST_PROTOCOL)

        # # #
        print("=" * 100)
        #

        result = {"status": "ok"}
        return result

    def create_total_pop(self, dic):
        print("900443-155 PotentialAlgo create_total_pop: \n", dic, "\n", "="*50)
        app_ = dic["app"]
        year= dic["year"]
        group = dic["group"]
        variable = dic["variable"]
        # print(self.entity_name)
        model_name_ = dic["dimensions"]["time_dim"]["model"]
        model_time_dim = apps.get_model(app_label=app_, model_name=model_name_)

        model_name_ = dic["dimensions"][self.entity_name+"_dim"]["model"]
        model_entity_dim = apps.get_model(app_label=app_, model_name=model_name_)

        model_name_ = "measuregroupdim"
        model_group_measure_dim = apps.get_model(app_label=app_, model_name=model_name_)

        model_name_ = dic["dimensions"]["measure_dim"]["model"]
        model_measure_dim = apps.get_model(app_label=app_, model_name=model_name_)

        model_name_ = dic["fact_model"]
        model_fact = apps.get_model(app_label=app_, model_name=model_name_)

        t = model_time_dim.objects.get(id=year)
        g, is_created = model_group_measure_dim.objects.get_or_create(group_name=group)
        m, is_created = model_measure_dim.objects.get_or_create(measure_group_dim=g, measure_name=variable)

        qs = model_entity_dim.objects.all()
        for e in qs:
            # print(e)
            s = 'model_fact.objects.get_or_create(time_dim=t, ' + self.entity_name + '_dim=e, measure_dim=m)'
            # print(s)
            ef, is_created = eval(s)
            ef.amount=1
            ef.save()

        result = {"status": "ok"}
        return result

    def get_similarity(self, dic):
        # print("s="*50)
        f = dic["f"]
        dependent_group = dic["dependent_group"]
        df_d = dic["df_d"]
        steph = dic["steph"]
        stepl = dic["stepl"]
        ll_dfs = dic["ll_dfs"]
        f_groups = dic["groups"]
        groups = f_groups[steph][stepl]["groups"]
        dff = pd.DataFrame(df_d.loc[:, f].astype(float))
        df_f = dff.copy()

        max_cut = groups[dependent_group][f]["max_cut"]
        min_cut = groups[dependent_group][f]["min_cut"]

        # max_cut = f_groups["max_cut"]
        # min_cut = f_groups["min_cut"]

        df_f = df_f.apply(lambda x: (x - min_cut) / (max_cut - min_cut))
        df_f[df_f < 0] = 0
        df_f[df_f > 1] = 1
        dffn = df_f.copy()
        llg_ = []
        for g in groups:
            dfg = ll_dfs[g].copy()
            lls = []
            if len(groups[g]) == 0:
                return 0
            for f_g in groups[g]:
                max_cut = groups[g][f_g]["max_cut"]
                min_cut = groups[g][f_g]["min_cut"]
                dfg_f = dfg[f_g]
                dfg_f = dfg_f.apply(lambda x: (x - min_cut) / (max_cut - min_cut))
                dfg_f[dfg_f < 0] = 0
                dfg_f[dfg_f > 1] = 1
                dfm = pd.merge(left=dffn, how='outer', right=dfg_f, left_index=True, right_index=True)
                dfm['diffd'] = dfm.apply(lambda x: abs(x[dfm.columns[0]] - x[dfm.columns[1]]), axis=1)
                dfm['diffr'] = dfm.apply(lambda x: abs(x[dfm.columns[0]] - (1-x[dfm.columns[1]])), axis=1)
                s_d = dfm['diffd'].mean()
                s_r = dfm['diffr'].mean()
                if s_d < s_r:
                    s = (1 - s_d)-0.7
                else:
                    s = (1 - s_r)-0.7
                if s < 0:
                    s = 0
                lls.append(s)
            llg_.append(mean(lls))
        return mean(llg_)

    def update_min_max_cuts(self, dic):
        print("90055-156 PotentialAlgo update_min_max_cuts: \n", dic, "\n", "="*50)
        app_ = dic["app"]
        year_ = str(dic["time_dim_value"])
        f = str(dic["var_obj"])
        # dependent_group = str(dic["dependent_group"])
        min_max_model_name_ = dic["min_max_model"]
        model_min_max = apps.get_model(app_label=app_, model_name=min_max_model_name_)
        # model_measure_dim = apps.get_model(app_label=app_, model_name="measuredim")
        file_path = os.path.join(self.PICKLE_PATH, "best_cut_"+f+"_"+year_+".pkl")
        print(file_path)
        with open(file_path, 'rb') as handle:
            best_cut = pickle.load(handle)
        # hh = best_cut["hh"]
        # ll = best_cut["ll"]
        # steph = best_cut["steph"]
        # stepl = best_cut["stepl"]
        file_path = os.path.join(self.PICKLE_PATH, "result_"+f+"_"+year_+".pkl")
        # print(file_path)
        with open(file_path, 'rb') as handle:
            results = pickle.load(handle)
        # print(results[hh][ll][int(f)])

        model_min_max.objects.filter(time_dim_id=year_).delete()

        # print("1111 best_cut\n", best_cut)
        f_groups = best_cut['f_groups']
        for g in f_groups:
            gfs = f_groups[g]
            # print("-"*60, "\n", g, "\n", gfs)
            for f in gfs:
                # print("-"*10, "\n", f, "\n", gfs[f])
                obj = model_min_max.objects.create(time_dim_id=year_, measure_dim_id=f,
                                                   min=round(100*gfs[f]["min_cut"])/100,
                                                   max=round(100*gfs[f]["max_cut"])/100)
        result = {"status": "ok"}
        return result

    def create_similarity(self, group_d, group, df_n1_all, df_n2_all, df_n1_, df_n2_, sign_n1, sign_n2,
                          similarity_n1, similarity_n2):

        # print("1 create_similarity df_n2_all\n", group, "\n", df_n2_all, df_n2_all.shape)

        df_n1_all = pd.merge(left=df_n1_all, how='outer', right=df_n1_, left_index=True, right_index=True)
        df_n2_all = pd.merge(left=df_n2_all, how='outer', right=df_n2_, left_index=True, right_index=True)

        # print("2 create_similarity df_n2_all\n", group, "\n", df_n2_all, df_n2_all.shape)

        # print(df_n2_all.head(100))
        for n in ["1", "2"]:
            # print(n)
            ll = []
            lls = []
            for k in self.options:
                # print(k[0], k[1])
                s_ = "abs(df_n" + n + "_all['" + k[0] + "-" + group_d + "'] - df_n" + n + "_all['" + k[1] + "-" + group + "'])"
                # print(s_)
                df_d = eval(s_)
                df_d_na=df_d.dropna()
                s_ = title='sim-n' + n + group + k
                df_d_na_ = self.add_entity_to_df1(df_d_na)
                self.add_to_save_all(title=s_, a=df_d_na_, cols=-1)
                s_d = df_d.sum()
                # print("s_d", s_d, df_d.mean())

                s_ = "abs(df_n" + n + "_all['" + k[0] + "-" + group_d + "'] - 1 + df_n" + n + "_all['" + k[1] + "-" + group + "'])"
                # print(s_)
                df_r = eval(s_)
                s_r = df_r.sum()
                # print("s_r", s_r, df_r.mean())

                if s_d < s_r:
                    d_ = s_d
                    s_ = "df_n" + n + "_all['" + group_d + '-' + group + '-' + k + "'] = df_d"
                    # print(s_)
                    exec(s_)
                    lls.append(1 - df_d.mean())
                    ll.append(1)
                else:
                    d_ = s_r
                    s_ = "df_n" + n + "_all['" + group_d + '-' + group + '-' + k + "'] = df_r"
                    # print(s_)
                    exec(s_)
                    lls.append(1 - df_r.mean())
                    ll.append(-1)
                # print("="*50)

            # print("sign_n" + n + ".loc[group] = ll", ll)
            exec("sign_n" + n + ".loc[group] = ll")
            # print("similarity_n" + n + ".loc[group] = lls", lls)
            exec("similarity_n" + n + ".loc[group] = lls")

        return group_d, group, df_n1_all, df_n2_all, df_n1_, df_n2_, sign_n1, sign_n2, similarity_n1, similarity_n2

    def add_entity_to_df(self, df, cols=None):
        df_ = df.copy()
        if cols is None:
            cols = [self.entity_name+'_name']
            df_c = df.columns
            for j in df_c:
                k = str(self.measures_name[self.measures_name["id"] == j]["measure_name"]).split("    ")[1].split("\n")[
                    0]
                cols.append(k)
        # print("add_entity_to_df", 1)
        df_ = df_.reset_index()
        # print("-"*10, "\nAA\n", "\n", df_)
        df_ = df_.merge(self.entities_name, how='inner', left_on=self.entity_name+'_dim', right_on='id').drop(
            [self.entity_name+'_dim', 'id'], axis=1)
        c_ = df_.pop(self.entity_name+'_name')
        df_.insert(0, self.entity_name+'_name', c_)
        # print("CC\n", "\n", df_)
        # print("CC1\n", df_.columns, "\n", cols)
        if isinstance(cols, pd.core.indexes.base.Index) or cols != -1:
            df_.columns = cols
        return df_

    def add_entity_to_df1(self, df, cols=None):
        df_ = df.copy()
        df_ = df_.reset_index()
        df_.columns=["id", "amount"]
        # print(1000, "\n", "-"*10, "\nAA\n", "\n", df_, self.entities_name)
        df_ = self.entities_name.merge(df_, how='inner', left_on='id', right_on='id')\
            # .drop(
            # [self.entity_name+'_dim', 'id'], axis=1)
        # print("BB3\n", "\n", df_)
        return df_

    # It seems that I can delete this function
    def save_to_excel(self, df, folder):
        df2 = df.copy()
        # print(self.save_to_file + ' -Before sleep save_to_excel- ' + folder)
        total, used, free = shutil.disk_usage("/")
        # print(' total: ' + str(total) + ' used: ' + str(used) + ' free: ' + str(free))
        nnn = 0
        try:
            with pd.ExcelWriter(self.save_to_file, engine='openpyxl', mode='a') as writer_:
                df2.to_excel(writer_, sheet_name=folder)
                writer_.save()
                time.sleep(5)
            if self.second_time_save != '':
                print("save ok:", self.second_time_save)
            self.second_time_save = ''
            nnn = 1
        except Exception as ee:
            print(ee)
            time.sleep(5)
            self.save_to_excel(df2, folder)
            self.second_time_save = self.save_to_file
            nnn = 1
        finally:
            if nnn == 0:
                print(self.save_to_file + ' 55 finally -' + str(nnn) + ' - ' + folder)
                time.sleep(5)
                print(self.save_to_file + ' 551 finally -' + str(nnn) + ' - ' + folder)
                self.save_to_excel(df2, folder)
                self.second_time_save = self.save_to_file

    def save_to_excel_(self):
        wb2 = Workbook()
        wb2.save(self.save_to_file)
        wb2.close()
        wb2 = None
        # print("save_to_excel_", self.save_to_file)

        with pd.ExcelWriter(self.save_to_file, engine='openpyxl', mode="a") as writer_o:
            for d in self.to_save:
                try:
                    # print("d[0]\n", d[0])
                    # print("d[1]\n", d[1])
                    d[0].to_excel(writer_o, sheet_name=d[1])
                except Exception as ex:
                    print("9006-3 " + str(ex))
            writer_o.save()
        wb = load_workbook(filename=self.save_to_file, read_only=False)
        del wb['Sheet']
        wb.save(self.save_to_file)
        wb.close()

    def save_to_excel_all_(self, year):
        save_to_file_all = os.path.join(self.TO_EXCEL_OUTPUT, "all_" + str(year) + ".xlsx")
        # print("save_to_excel_all_", save_to_file_all)
        wb2 = Workbook()
        wb2.save(save_to_file_all)
        wb2.close()

        with pd.ExcelWriter(save_to_file_all, engine='openpyxl', mode="a") as writer:
            for d in self.to_save_all:
                try:
                    d[0].to_excel(writer, sheet_name=d[1])
                except Exception as ex:
                    print("9006-3 " + d[2] + str(ex))
            writer.save()

        # wb = load_workbook(filename=save_to_file_all, read_only=False)
        # del wb['Sheet']
        # wb.save(self.save_to_file)
        # wb.close()

    def add_to_save(self, title, a, cols):
        ai = pd.DataFrame(a)
        ai.index = self.df_index
        df_ = self.add_entity_to_df(ai, cols=cols)
        df_.dropna(how='all', axis=1, inplace=True)
        self.to_save.append((df_.copy(), title))

    def add_to_save_all(self, title, a, cols):
        ai = pd.DataFrame(a)
        ai.dropna(how='all', axis=1, inplace=True)
        self.to_save_all.append((ai.copy(), title))

    def clean_rows(self, a, j, side="L"):
        # print("j", j)
        # a.sort(axis=1)
        # if j == 1:
        # self.add_to_save(title='Sort '+side+"-"+str(j), a=a, cols=-1)
        #
        a_ = pd.DataFrame(a.copy())
        a_ = a_.loc[a_.apply(lambda x: x.count(), axis=1) > 4]
        a_ = a_.to_numpy()
        # print(a.shape, "\n", a_.shape)
        if a_.shape[0] == 0:
            return a
        d = a[:, j:1 + j] - a[:, 0:1]
        d_ = a_[:, j:1 + j] - a_[:, 0:1]
        d_m = np.nanmean(d_, axis=0)
        # print(j, "\n", d_m, "\n", d_m[0])
        if d_m[0] < 0.03:
            #
            df_d = pd.DataFrame(d)
            df_d.index = self.df_index
            df_ = self.add_entity_to_df(df_d, cols=-1)
            self.to_save.append((df_.copy(), 'D_' + side + ' - ' + str(j)))
            # print("d=\n", d)
            #
            b = pd.DataFrame(a.copy())
            b.loc[b.apply(lambda x: x.count(), axis=1) > 4, [j]] = np.nan
            #
            b.index = self.df_index
            df_ = self.add_entity_to_df(b, cols=-1)
            self.to_save.append((df_.copy(), side + ' b ' + str(j)))
            #
            b = b.to_numpy()
            if j + 1 < b.shape[1]:
                b = self.clean_rows(b, j + 1, side)
        else:
            b = a.copy()
        return b

    def move_elements_to_right(self, row):
        n_nna = row.count()
        n_na = row.isna().sum().sum()
        n = n_na + n_nna
        # print("line 1 row=", n, n_na, n_nna)
        # print(row)
        row_c = row.copy()
        for j in range(n-1):
            try:
                if str(row_c[n-j-1:n-j].iloc[0]) == "nan":
                    for z in range(0, n-1-j):
                        if str(row_c[n-j-2-z:n-j-1-z].iloc[0]) != "nan":
                            row_c[n-j-1:n-j].iloc[0] = row_c[n-j-2-z:n-j-1-z].iloc[0]
                            row_c[n-j-2-z:n-j-1-z].iloc[0] = np.nan
                            break
            except Exception as ex:
                print("ex1: "+str(ex))
        # print("-"*100)
        # print(row)
        # print(row_c)
        # print("-"*100)
        return row_c

    def revers_elements_in_row(self, row):
        n_nna = row.count()
        n_na = row.isna().sum().sum()
        n = n_na + n_nna
        row_ = row.copy()
        for j in range(n):
            row_[j:j+1] = row[n-j-1:n-j]
        return row_

    def twenty_rule(self, row):
        # print("-1"*10, "\nrow:\n", row, "-2"*10, "\n")
        n_row = row.count()
        # print("n_row= ", n_row)
        if n_row < 5:
            n = 0
        elif n_row == 5:
            n = 1
        else:
            n = math.ceil(n_row * 0.2)
        min = n_row
        row_best = row
        if n == 0:
            return row_best

        # print("-4"*10)
        # print("n= ", n, "\n range(n+1)=", range(n+1))
        # print("-5"*10)
        for j in range(n + 1):
            row_c = row.copy()
            row_c[:j] = np.nan
            # print("-6"*10)
            # print(row_c)
            # print("-7"*10)
            row_c[n_row - (n - j):] = np.nan
            # print(row_c)
            # print("-8"*10)
            # print("max=", row_c.max(), "min=", row_c.min())
            # print("-9"*10)
            if row_c.max() - row_c.min() < min:
                min = row_c.max() - row_c.min()
                row_best = row_c.copy()
        return row_best

    def thirty_rule(self, row):
        n_row = row.count()
        row_best = row
        if (n_row > 4) and abs(row.max() - row.min()) >= 0.3:
            n = 1
            min = n_row
            for j in range(n + 1):
                row_c = row.copy()
                row_c[:j] = np.nan
                row_c[n_row - (n - j):] = np.nan
                if (row_c.max() - row_c.min()) < min:
                    min = row_c.max() - row_c.min()
                    row_best = row_c.copy()
        if abs(row_best.max() - row_best.min()) >= 0.3:
            row_best[:] = np.nan
        return row_best

    def min_max_rule(self, row):
        row_ = row[0:2].copy()
        row_[:] = np.nan
        row_[0] = row.min()
        row_[1] = row.max()
        return row_